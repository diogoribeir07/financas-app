import os
import uuid
from datetime import datetime
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_EXCEL_PATH = os.path.join(BASE_DIR, "Gestao_Financeira.xlsx")

DEFAULT_ACCOUNTS = [
    {
        "id": "geral",
        "name": "Dinheiro Geral",
        "icon": "DG",
        "initial_balance": 370.43,
        "color": "#71717A",
        "description": "Gastos do dia-a-dia e dinheiro geral (Cofrinho)",
    },
    {
        "id": "comida_mae",
        "name": "Comida Mãe",
        "icon": "CM",
        "initial_balance": 0.0,
        "color": "#52525B",
        "description": "Gastos de comida e supermercado (Mãe)",
    },
    {
        "id": "comida_pai",
        "name": "Comida Pai",
        "icon": "CP",
        "initial_balance": 56.68,
        "color": "#52525B",
        "description": "Gastos de comida e supermercado (Pai)",
    },
    {
        "id": "bolsa_volumes",
        "name": "Bolsa & Grandes Volumes",
        "icon": "BV",
        "initial_balance": 0.0,
        "color": "#3F3F46",
        "description": "Bolsa, poupanças e grandes volumes",
    },
]

DEFAULT_CATEGORIES = [
    {"id": "alimentacao", "name": "Alimentação", "icon": "AL", "type": "Despesa", "budget": 250.0, "color": "#71717A", "chips": "Café,Almoço,Jantar,Padaria,Uber Eats"},
    {"id": "supermercado", "name": "Supermercado", "icon": "SM", "type": "Despesa", "budget": 300.0, "color": "#71717A", "chips": "Continente,Pingo Doce,Lidl,Mercadona,Talho"},
    {"id": "transportes", "name": "Transportes", "icon": "TR", "type": "Despesa", "budget": 150.0, "color": "#71717A", "chips": "Gasolina,Portagem,Uber / Bolt,Passe,Estacionamento"},
    {"id": "habitacao", "name": "Habitação & Contas", "icon": "HC", "type": "Despesa", "budget": 600.0, "color": "#71717A", "chips": "Renda,Eletricidade,Água,Internet"},
    {"id": "lazer", "name": "Lazer & Saídas", "icon": "LZ", "type": "Despesa", "budget": 150.0, "color": "#71717A", "chips": "Cinema,Jantar Fora,Copos,Viagem,Desporto"},
    {"id": "saude", "name": "Saúde & Bem-estar", "icon": "SB", "type": "Despesa", "budget": 80.0, "color": "#71717A", "chips": "Farmácia,Consulta,Ginásio,Barbeiro"},
    {"id": "subscricoes", "name": "Subscrições", "icon": "SU", "type": "Despesa", "budget": 50.0, "color": "#71717A", "chips": "Spotify,Netflix,iCloud,ChatGPT,Telemóvel"},
    {"id": "compras", "name": "Compras Pessoais", "icon": "CP", "type": "Despesa", "budget": 120.0, "color": "#71717A", "chips": "Roupa,Calçado,Amazon,Tecnologia,Presente"},
    {"id": "outros", "name": "Outros", "icon": "OU", "type": "Despesa", "budget": 100.0, "color": "#71717A", "chips": "Multibanco,Imprevisto,Diversos"},
    {"id": "bolsa", "name": "Bolsa & Apoios", "icon": "BA", "type": "Receita", "budget": 0.0, "color": "#71717A", "chips": "Bolsa Mensal,Apoio,Prémio"},
    {"id": "reforco_pais", "name": "Reforço Mãe & Pai", "icon": "RP", "type": "Receita", "budget": 0.0, "color": "#71717A", "chips": "Transferência Mãe,Transferência Pai,Mesada Comida"},
    {"id": "salario", "name": "Salário & Rendimento", "icon": "SR", "type": "Receita", "budget": 0.0, "color": "#71717A", "chips": "Salário,Reembolso MBWay,Trabalho Extra"},
]


class ExcelManager:
    def __init__(self, excel_path: str = DEFAULT_EXCEL_PATH):
        self.excel_path = excel_path
        self.ensure_workbook()

    def ensure_workbook(self) -> None:
        """Cria ou atualiza o ficheiro Excel com suporte a Múltiplas Contas, Categorias Editáveis e Recorrentes."""
        if os.path.exists(self.excel_path):
            try:
                wb = openpyxl.load_workbook(self.excel_path)
                has_legacy_combined_food = False
                if "Contas" in wb.sheetnames:
                    for r in range(2, wb["Contas"].max_row + 1):
                        if str(wb["Contas"].cell(row=r, column=1).value) == "comida_pais":
                            has_legacy_combined_food = True
                            break
                needs_upgrade = (
                    "Contas" not in wb.sheetnames
                    or "Recorrentes" not in wb.sheetnames
                    or wb["Transacoes"].cell(row=1, column=5).value != "Conta"
                    or has_legacy_combined_food
                )
                wb.close()
                if not needs_upgrade:
                    return
                # Preservar transações existentes antes de atualizar a estrutura
                old_txs = self._extract_legacy_transactions()
                os.remove(self.excel_path)
                self._create_fresh_workbook()
                if old_txs:
                    self.sync_transactions_batch(old_txs)
                return
            except Exception:
                pass

        self._create_fresh_workbook()

    def _extract_legacy_transactions(self) -> List[Dict[str, Any]]:
        try:
            wb = openpyxl.load_workbook(self.excel_path, data_only=True)
            if "Transacoes" not in wb.sheetnames:
                wb.close()
                return []
            ws = wb["Transacoes"]
            headers = [str(ws.cell(row=1, column=i).value or "") for i in range(1, ws.max_column + 1)]
            has_account_col = "Conta" in headers
            txs = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not any(row):
                    continue
                if has_account_col:
                    txs.append({
                        "id": str(row[0]),
                        "date": str(row[1])[:10],
                        "time": str(row[2] or "12:00")[:5],
                        "type": str(row[3] or "Despesa"),
                        "account": str(row[4] or "Dinheiro Geral"),
                        "to_account": str(row[5] or ""),
                        "category": str(row[6] or "Outros"),
                        "description": str(row[7] or ""),
                        "amount": float(row[8] or 0.0),
                        "payment_method": str(row[9] or "MBWay"),
                        "month_year": str(row[10] or str(row[1])[:7]),
                    })
                else:
                    txs.append({
                        "id": str(row[0]),
                        "date": str(row[1])[:10],
                        "time": str(row[2] or "12:00")[:5],
                        "type": str(row[3] or "Despesa"),
                        "account": "Dinheiro Geral",
                        "to_account": "",
                        "category": str(row[4] or "Outros"),
                        "description": str(row[6] or ""),
                        "amount": float(row[7] or 0.0),
                        "payment_method": str(row[8] or "MBWay"),
                        "month_year": str(row[9] or str(row[1])[:7]),
                    })
            wb.close()
            return txs
        except Exception:
            return []

    def _create_fresh_workbook(self) -> None:
        wb = openpyxl.Workbook()
        default_sheet = wb.active

        ws_tx = wb.create_sheet(title="Transacoes")
        self._setup_transacoes_sheet(ws_tx)

        ws_acc = wb.create_sheet(title="Contas")
        self._setup_contas_sheet(ws_acc)

        ws_cat = wb.create_sheet(title="Categorias")
        self._setup_categorias_sheet(ws_cat)

        ws_rec = wb.create_sheet(title="Recorrentes")
        self._setup_recorrentes_sheet(ws_rec)

        ws_dash = wb.create_sheet(title="Dashboard", index=0)
        self._setup_dashboard_sheet(ws_dash)

        wb.remove(default_sheet)
        wb.save(self.excel_path)

    def _setup_transacoes_sheet(self, ws) -> None:
        headers = [
            "ID",
            "Data",
            "Hora",
            "Tipo",
            "Conta",
            "Conta Destino",
            "Categoria",
            "Descrição",
            "Valor (€)",
            "Método de Pagamento",
            "Mês_Ano",
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="medium", color="0F172A"),
        )

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:K1"

        col_widths = [20, 14, 10, 14, 24, 24, 22, 32, 16, 20, 12]
        for idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

    def _setup_contas_sheet(self, ws) -> None:
        headers = ["ID", "Nome da Conta", "Ícone", "Saldo Inicial (€)", "Cor Hex", "Descrição"]
        ws.append(headers)

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[1].height = 26
        ws.freeze_panes = "A2"

        data_font = Font(name="Segoe UI", size=10)
        for row_idx, acc in enumerate(DEFAULT_ACCOUNTS, start=2):
            ws.append([
                acc["id"],
                acc["name"],
                acc["icon"],
                acc["initial_balance"],
                acc["color"],
                acc["description"],
            ])
            for col_idx in range(1, 7):
                c = ws.cell(row=row_idx, column=col_idx)
                c.font = data_font
                if col_idx == 4:
                    c.number_format = '#,##0.00 "€"'
                    c.alignment = Alignment(horizontal="right", vertical="center")
                elif col_idx in (1, 3, 5):
                    c.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[row_idx].height = 22

        col_widths = [18, 26, 10, 20, 14, 38]
        for idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

    def _setup_categorias_sheet(self, ws) -> None:
        headers = ["ID", "Categoria", "Ícone", "Tipo", "Orçamento Mensal (€)", "Cor Hex", "Sugestões Rápidas"]
        ws.append(headers)

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[1].height = 26
        ws.freeze_panes = "A2"

        data_font = Font(name="Segoe UI", size=10)
        for row_idx, cat in enumerate(DEFAULT_CATEGORIES, start=2):
            ws.append([
                cat["id"],
                cat["name"],
                cat["icon"],
                cat["type"],
                cat["budget"],
                cat["color"],
                cat.get("chips", ""),
            ])
            for col_idx in range(1, 8):
                c = ws.cell(row=row_idx, column=col_idx)
                c.font = data_font
                if col_idx == 5:
                    c.number_format = '#,##0.00 "€"'
                    c.alignment = Alignment(horizontal="right", vertical="center")
                elif col_idx in (1, 3, 4, 6):
                    c.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[row_idx].height = 22

        col_widths = [16, 26, 10, 14, 22, 14, 42]
        for idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

    def _setup_recorrentes_sheet(self, ws) -> None:
        headers = [
            "ID",
            "Descrição",
            "Tipo",
            "Categoria",
            "Conta",
            "Valor (€)",
            "Frequência",
            "Dia do Mês",
            "Método de Pagamento",
            "Último Mês Processado",
            "Ativo",
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[1].height = 26
        ws.freeze_panes = "A2"

        col_widths = [18, 28, 14, 22, 24, 16, 16, 14, 20, 22, 12]
        for idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

    def _setup_dashboard_sheet(self, ws) -> None:
        ws.views.sheetView[0].showGridLines = False

        # Banner de Título
        ws.merge_cells("B2:I3")
        title_cell = ws["B2"]
        title_cell.value = "PAINEL DE GESTÃO FINANCEIRA PESSOAL"
        title_cell.font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="18181B", end_color="18181B", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        # Cartões KPI (Linhas 5 e 6)
        kpis = [
            ("B", "C", "TOTAL DESPESAS", '=SUMIF(Transacoes!D2:D5000, "Despesa", Transacoes!I2:I5000)', "EF4444", '#,##0.00 "€"'),
            ("D", "E", "TOTAL RECEITAS", '=SUMIF(Transacoes!D2:D5000, "Receita", Transacoes!I2:I5000)', "10B981", '#,##0.00 "€"'),
            ("F", "G", "PATRIMÓNIO TOTAL", "=SUM(Contas!D2:D50)+D6-B6", "3B82F6", '#,##0.00 "€"'),
            ("H", "I", "Nº MOVIMENTOS", "=COUNTA(Transacoes!A2:A5000)", "6366F1", "#,##0"),
        ]

        for col_start, col_end, label, formula, color_hex, num_fmt in kpis:
            ws.merge_cells(f"{col_start}5:{col_end}5")
            ws.merge_cells(f"{col_start}6:{col_end}6")

            lbl_cell = ws[f"{col_start}5"]
            lbl_cell.value = label
            lbl_cell.font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
            lbl_cell.fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")
            lbl_cell.alignment = Alignment(horizontal="center", vertical="center")

            val_cell = ws[f"{col_start}6"]
            val_cell.value = formula
            val_cell.font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
            val_cell.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            val_cell.alignment = Alignment(horizontal="center", vertical="center")
            val_cell.number_format = num_fmt

        ws.row_dimensions[5].height = 20
        ws.row_dimensions[6].height = 32

        # Tabela de Resumo por Categoria (A partir da linha 9)
        ws["B8"].value = "Resumo de Gastos por Categoria vs. Orçamento"
        ws["B8"].font = Font(name="Segoe UI", size=12, bold=True, color="0F172A")

        table_headers = ["Categoria", "Ícone", "Orçamento (€)", "Total Gasto (€)", "Disponível (€)", "% Utilizado"]
        header_cols = ["B", "C", "D", "E", "F", "G"]
        for c_let, h_text in zip(header_cols, table_headers):
            c = ws[f"{c_let}9"]
            c.value = h_text
            c.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[9].height = 24

        expense_categories = [cat for cat in DEFAULT_CATEGORIES if cat["type"] == "Despesa"]
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        for idx, _ in enumerate(expense_categories):
            r = 10 + idx
            cat_row = idx + 2
            ws[f"B{r}"].value = f"=Categorias!B{cat_row}"
            ws[f"C{r}"].value = f"=Categorias!C{cat_row}"
            ws[f"D{r}"].value = f"=Categorias!E{cat_row}"
            ws[f"E{r}"].value = f'=SUMIFS(Transacoes!$I$2:$I$5000, Transacoes!$G$2:$G$5000, B{r}, Transacoes!$D$2:$D$5000, "Despesa")'
            ws[f"F{r}"].value = f"=D{r}-E{r}"
            ws[f"G{r}"].value = f"=IF(D{r}>0, E{r}/D{r}, 0)"

            fill_color = "FFFFFF" if idx % 2 == 0 else "F8FAFC"
            row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

            for c_let in header_cols:
                cell = ws[f"{c_let}{r}"]
                cell.font = Font(name="Segoe UI", size=10, color="1E293B")
                cell.fill = row_fill
                cell.border = thin_border
                if c_let in ("D", "E", "F"):
                    cell.number_format = '#,##0.00 "€"'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif c_let == "G":
                    cell.number_format = "0.0%"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif c_let == "C":
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[r].height = 22

        last_cat_row = 10 + len(expense_categories) - 1
        total_row = last_cat_row + 1
        ws[f"B{total_row}"].value = "TOTAL"
        ws[f"C{total_row}"].value = ""
        ws[f"D{total_row}"].value = f"=SUM(D10:D{last_cat_row})"
        ws[f"E{total_row}"].value = f"=SUM(E10:E{last_cat_row})"
        ws[f"F{total_row}"].value = f"=D{total_row}-E{total_row}"
        ws[f"G{total_row}"].value = f"=IF(D{total_row}>0, E{total_row}/D{total_row}, 0)"

        for c_let in header_cols:
            cell = ws[f"{c_let}{total_row}"]
            cell.font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
            cell.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
            if c_let in ("D", "E", "F"):
                cell.number_format = '#,##0.00 "€"'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c_let == "G":
                cell.number_format = "0.0%"
                cell.alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[total_row].height = 24

        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Gastos vs. Orçamento por Categoria (€)"
        chart.y_axis.title = "Valor (€)"
        chart.x_axis.title = "Categoria"
        data_ref = Reference(ws, min_col=4, min_row=9, max_col=5, max_row=last_cat_row)
        cats_ref = Reference(ws, min_col=2, min_row=10, max_row=last_cat_row)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.width = 16
        chart.height = 10
        ws.add_chart(chart, f"B{total_row + 2}")

        ws.column_dimensions["A"].width = 4
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 10
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 18
        ws.column_dimensions["F"].width = 18
        ws.column_dimensions["G"].width = 16
        ws.column_dimensions["H"].width = 16
        ws.column_dimensions["I"].width = 16

    def _style_transaction_row(self, ws, row_idx: int, tx_type: str) -> None:
        fill_color = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="F1F5F9"),
            right=Side(style="thin", color="F1F5F9"),
            top=Side(style="thin", color="F1F5F9"),
            bottom=Side(style="thin", color="E2E8F0"),
        )
        base_font = Font(name="Segoe UI", size=10, color="1E293B")
        if tx_type == "Receita":
            val_color = "16A34A"
        elif tx_type == "Transferência":
            val_color = "2563EB"
        else:
            val_color = "0F172A"

        for col_idx in range(1, 12):
            c = ws.cell(row=row_idx, column=col_idx)
            c.fill = row_fill
            c.border = thin_border
            c.font = base_font
            if col_idx == 9:
                c.number_format = '#,##0.00 "€"'
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.font = Font(name="Segoe UI", size=10, bold=True, color=val_color)
            elif col_idx in (1, 2, 3, 4, 10, 11):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[row_idx].height = 22

    # =========================================================================
    # CONTAS (MÚLTIPLAS CONTAS & SALDOS)
    # =========================================================================
    def get_accounts(self) -> List[Dict[str, Any]]:
        self.ensure_workbook()
        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        ws = wb["Contas"]
        accounts = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[1]:
                continue
            accounts.append({
                "id": str(row[0] or uuid.uuid4().hex[:8]),
                "name": str(row[1]),
                "icon": str(row[2] or "💳"),
                "initial_balance": round(float(row[3] or 0.0), 2),
                "color": str(row[4] or "#6366F1"),
                "description": str(row[5] or ""),
            })
        wb.close()
        if not accounts:
            accounts = [dict(a) for a in DEFAULT_ACCOUNTS]

        # Calcular saldo atual de cada conta com base nas transações
        txs = self.get_all_transactions()
        for acc in accounts:
            balance = acc["initial_balance"]
            acc_name = acc["name"]
            for t in txs:
                amt = float(t.get("amount", 0.0))
                t_type = t.get("type", "Despesa")
                t_acc = t.get("account", "Dinheiro Geral")
                t_to_acc = t.get("to_account", "")

                if t_type == "Receita" and t_acc == acc_name:
                    balance += amt
                elif t_type == "Despesa" and t_acc == acc_name:
                    balance -= amt
                elif t_type == "Transferência":
                    if t_acc == acc_name:
                        balance -= amt
                    if t_to_acc == acc_name:
                        balance += amt
            acc["current_balance"] = round(balance, 2)

        return accounts

    def save_account(self, acc_data: Dict[str, Any]) -> Dict[str, Any]:
        self.ensure_workbook()
        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb["Contas"]

        acc_id = str(acc_data.get("id") or f"acc_{uuid.uuid4().hex[:6]}")
        name = str(acc_data.get("name", "Nova Conta")).strip()
        icon = str(acc_data.get("icon") or "💳").strip()
        initial_balance = round(float(acc_data.get("initial_balance") or 0.0), 2)
        color = str(acc_data.get("color") or "#6366F1")
        description = str(acc_data.get("description") or "")

        found_row = None
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(row=r, column=1).value) == acc_id or str(ws.cell(row=r, column=2).value) == name:
                found_row = r
                break

        if found_row:
            ws.cell(row=found_row, column=1, value=acc_id)
            ws.cell(row=found_row, column=2, value=name)
            ws.cell(row=found_row, column=3, value=icon)
            ws.cell(row=found_row, column=4, value=initial_balance)
            ws.cell(row=found_row, column=5, value=color)
            ws.cell(row=found_row, column=6, value=description)
        else:
            ws.append([acc_id, name, icon, initial_balance, color, description])
            r = ws.max_row
            ws.cell(row=r, column=4).number_format = '#,##0.00 "€"'

        wb.save(self.excel_path)
        wb.close()
        return {
            "id": acc_id,
            "name": name,
            "icon": icon,
            "initial_balance": initial_balance,
            "color": color,
            "description": description,
        }

    def delete_account(self, acc_id: str) -> bool:
        self.ensure_workbook()
        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb["Contas"]
        deleted = False
        for r in range(ws.max_row, 1, -1):
            if str(ws.cell(row=r, column=1).value) == str(acc_id):
                ws.delete_rows(r, 1)
                deleted = True
                break
        if deleted:
            wb.save(self.excel_path)
        wb.close()
        return deleted

    # =========================================================================
    # CATEGORIAS / TIPOS DE GASTO EDITÁVEIS
    # =========================================================================
    def get_categories(self) -> List[Dict[str, Any]]:
        self.ensure_workbook()
        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        ws = wb["Categorias"]
        cats = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[1]:
                continue
            chips_raw = str(row[6] or "") if len(row) > 6 else ""
            chips_list = [c.strip() for c in chips_raw.split(",") if c.strip()]
            cats.append({
                "id": str(row[0] or row[1].lower()),
                "name": str(row[1]),
                "icon": str(row[2] or "📌"),
                "type": str(row[3] or "Despesa"),
                "budget": round(float(row[4] or 0.0), 2),
                "color": str(row[5] or "#64748B"),
                "chips": chips_list if chips_list else ["Registo Rápido"],
            })
        wb.close()
        return cats if cats else [dict(c) for c in DEFAULT_CATEGORIES]

    def save_category(self, cat_data: Dict[str, Any]) -> Dict[str, Any]:
        self.ensure_workbook()
        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb["Categorias"]

        cat_id = str(cat_data.get("id") or f"cat_{uuid.uuid4().hex[:6]}")
        name = str(cat_data.get("name", "Nova Categoria")).strip()
        icon = str(cat_data.get("icon") or "📌").strip()
        c_type = str(cat_data.get("type") or "Despesa")
        budget = round(float(cat_data.get("budget") or 0.0), 2)
        color = str(cat_data.get("color") or "#6366F1")
        chips = cat_data.get("chips", [])
        if isinstance(chips, list):
            chips_str = ",".join(chips)
        else:
            chips_str = str(chips or "")

        found_row = None
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(row=r, column=1).value) == cat_id:
                found_row = r
                break

        if found_row:
            old_name = str(ws.cell(row=found_row, column=2).value)
            ws.cell(row=found_row, column=2, value=name)
            ws.cell(row=found_row, column=3, value=icon)
            ws.cell(row=found_row, column=4, value=c_type)
            ws.cell(row=found_row, column=5, value=budget)
            ws.cell(row=found_row, column=6, value=color)
            ws.cell(row=found_row, column=7, value=chips_str)
            # Se o nome da categoria mudou, atualizar nas transações existentes também!
            if old_name and old_name != name and "Transacoes" in wb.sheetnames:
                ws_tx = wb["Transacoes"]
                for tx_r in range(2, ws_tx.max_row + 1):
                    if str(ws_tx.cell(row=tx_r, column=7).value) == old_name:
                        ws_tx.cell(row=tx_r, column=7, value=name)
        else:
            ws.append([cat_id, name, icon, c_type, budget, color, chips_str])
            ws.cell(row=ws.max_row, column=5).number_format = '#,##0.00 "€"'

        wb.save(self.excel_path)
        wb.close()
        return {
            "id": cat_id,
            "name": name,
            "icon": icon,
            "type": c_type,
            "budget": budget,
            "color": color,
            "chips": [c.strip() for c in chips_str.split(",") if c.strip()],
        }

    def delete_category(self, cat_id: str) -> bool:
        self.ensure_workbook()
        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb["Categorias"]
        deleted = False
        for r in range(ws.max_row, 1, -1):
            if str(ws.cell(row=r, column=1).value) == str(cat_id):
                ws.delete_rows(r, 1)
                deleted = True
                break
        if deleted:
            wb.save(self.excel_path)
        wb.close()
        return deleted

    # =========================================================================
    # TRANSAÇÕES & TRANSFERÊNCIAS ENTRE CONTAS
    # =========================================================================
    def get_all_transactions(self) -> List[Dict[str, Any]]:
        self.ensure_workbook()
        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        ws = wb["Transacoes"]
        transactions = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue
            tx_id = str(row[0]) if row[0] else f"TX-{uuid.uuid4().hex[:8].upper()}"
            date_val = row[1]
            if isinstance(date_val, datetime):
                date_str = date_val.strftime("%Y-%m-%d")
            else:
                date_str = str(date_val or datetime.now().strftime("%Y-%m-%d"))[:10]

            time_val = str(row[2] or "12:00")[:5]
            tx_type = str(row[3] or "Despesa")
            account = str(row[4] or "Dinheiro Geral")
            to_account = str(row[5] or "")
            category = str(row[6] or "Outros")
            description = str(row[7] or "")
            try:
                amount = float(row[8]) if row[8] is not None else 0.0
            except (ValueError, TypeError):
                amount = 0.0
            payment_method = str(row[9] or "MBWay")
            month_year = str(row[10] or date_str[:7])

            transactions.append({
                "id": tx_id,
                "date": date_str,
                "time": time_val,
                "type": tx_type,
                "account": account,
                "to_account": to_account,
                "category": category,
                "description": description,
                "amount": round(amount, 2),
                "payment_method": payment_method,
                "month_year": month_year,
            })

        wb.close()
        transactions.sort(key=lambda x: (x["date"], x["time"], x["id"]), reverse=True)
        return transactions

    def add_transaction(self, tx: Dict[str, Any]) -> Dict[str, Any]:
        self.ensure_workbook()
        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb["Transacoes"]

        tx_id = tx.get("id") or f"TX-{datetime.now().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:4].upper()}"
        date_str = tx.get("date") or datetime.now().strftime("%Y-%m-%d")
        time_str = tx.get("time") or datetime.now().strftime("%H:%M")
        tx_type = tx.get("type") or "Despesa"
        account = tx.get("account") or "Dinheiro Geral"
        to_account = tx.get("to_account") or ""
        category = tx.get("category") or ("Transferência" if tx_type == "Transferência" else "Outros")
        description = tx.get("description") or ""
        amount = round(float(tx.get("amount", 0.0)), 2)
        payment_method = tx.get("payment_method") or "MBWay"
        month_year = tx.get("month_year") or date_str[:7]

        for row_idx in range(2, ws.max_row + 1):
            if str(ws.cell(row=row_idx, column=1).value) == str(tx_id):
                wb.close()
                return {
                    "id": tx_id,
                    "date": date_str,
                    "time": time_str,
                    "type": tx_type,
                    "account": account,
                    "to_account": to_account,
                    "category": category,
                    "description": description,
                    "amount": amount,
                    "payment_method": payment_method,
                    "month_year": month_year,
                }

        row_values = [
            tx_id,
            date_str,
            time_str,
            tx_type,
            account,
            to_account,
            category,
            description,
            amount,
            payment_method,
            month_year,
        ]
        ws.append(row_values)
        self._style_transaction_row(ws, ws.max_row, tx_type)

        wb.save(self.excel_path)
        wb.close()

        return {
            "id": tx_id,
            "date": date_str,
            "time": time_str,
            "type": tx_type,
            "account": account,
            "to_account": to_account,
            "category": category,
            "description": description,
            "amount": amount,
            "payment_method": payment_method,
            "month_year": month_year,
        }

    def update_transaction(self, tx_id: str, tx_data: Dict[str, Any]) -> Optional[Dict[str, Any]]:
        self.ensure_workbook()
        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb["Transacoes"]
        updated = None

        for r in range(2, ws.max_row + 1):
            if str(ws.cell(row=r, column=1).value) == str(tx_id):
                date_str = tx_data.get("date") or str(ws.cell(row=r, column=2).value)[:10]
                time_str = tx_data.get("time") or str(ws.cell(row=r, column=3).value or "12:00")[:5]
                tx_type = tx_data.get("type") or str(ws.cell(row=r, column=4).value or "Despesa")
                account = tx_data.get("account") or str(ws.cell(row=r, column=5).value or "Dinheiro Geral")
                to_account = tx_data.get("to_account") if "to_account" in tx_data else str(ws.cell(row=r, column=6).value or "")
                category = tx_data.get("category") or str(ws.cell(row=r, column=7).value or "Outros")
                description = tx_data.get("description") if "description" in tx_data else str(ws.cell(row=r, column=8).value or "")
                amount = round(float(tx_data.get("amount", ws.cell(row=r, column=9).value or 0.0)), 2)
                payment_method = tx_data.get("payment_method") or str(ws.cell(row=r, column=10).value or "MBWay")
                month_year = date_str[:7]

                ws.cell(row=r, column=2, value=date_str)
                ws.cell(row=r, column=3, value=time_str)
                ws.cell(row=r, column=4, value=tx_type)
                ws.cell(row=r, column=5, value=account)
                ws.cell(row=r, column=6, value=to_account)
                ws.cell(row=r, column=7, value=category)
                ws.cell(row=r, column=8, value=description)
                ws.cell(row=r, column=9, value=amount)
                ws.cell(row=r, column=10, value=payment_method)
                ws.cell(row=r, column=11, value=month_year)
                self._style_transaction_row(ws, r, tx_type)

                updated = {
                    "id": tx_id,
                    "date": date_str,
                    "time": time_str,
                    "type": tx_type,
                    "account": account,
                    "to_account": to_account,
                    "category": category,
                    "description": description,
                    "amount": amount,
                    "payment_method": payment_method,
                    "month_year": month_year,
                }
                break

        if updated:
            wb.save(self.excel_path)
        wb.close()
        return updated

    def delete_transaction(self, tx_id: str) -> bool:
        self.ensure_workbook()
        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb["Transacoes"]
        deleted = False

        for row_idx in range(ws.max_row, 1, -1):
            if str(ws.cell(row=row_idx, column=1).value) == str(tx_id):
                ws.delete_rows(row_idx, 1)
                deleted = True
                break

        if deleted:
            wb.save(self.excel_path)
        wb.close()
        return deleted

    def sync_transactions_batch(self, incoming_txs: List[Dict[str, Any]]) -> Dict[str, int]:
        self.ensure_workbook()
        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb["Transacoes"]

        existing_ids = set()
        for row_idx in range(2, ws.max_row + 1):
            val = ws.cell(row=row_idx, column=1).value
            if val:
                existing_ids.add(str(val))

        added_count = 0
        for tx in incoming_txs:
            tx_id = str(tx.get("id", ""))
            if not tx_id or tx_id in existing_ids:
                continue

            date_str = tx.get("date") or datetime.now().strftime("%Y-%m-%d")
            time_str = tx.get("time") or datetime.now().strftime("%H:%M")
            tx_type = tx.get("type") or "Despesa"
            account = tx.get("account") or "Dinheiro Geral"
            to_account = tx.get("to_account") or ""
            category = tx.get("category") or "Outros"
            description = tx.get("description") or ""
            amount = round(float(tx.get("amount", 0.0)), 2)
            payment_method = tx.get("payment_method") or "MBWay"
            month_year = tx.get("month_year") or date_str[:7]

            ws.append([
                tx_id,
                date_str,
                time_str,
                tx_type,
                account,
                to_account,
                category,
                description,
                amount,
                payment_method,
                month_year,
            ])
            self._style_transaction_row(ws, ws.max_row, tx_type)
            existing_ids.add(tx_id)
            added_count += 1

        if added_count > 0:
            wb.save(self.excel_path)
        wb.close()
        return {"added": added_count, "total": len(existing_ids)}

    # =========================================================================
    # DEDUÇÕES / MOVIMENTOS RECORRENTES
    # =========================================================================
    def get_recurring(self) -> List[Dict[str, Any]]:
        self.ensure_workbook()
        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        ws = wb["Recorrentes"]
        items = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            items.append({
                "id": str(row[0]),
                "description": str(row[1] or ""),
                "type": str(row[2] or "Despesa"),
                "category": str(row[3] or "Subscrições"),
                "account": str(row[4] or "Dinheiro Geral"),
                "amount": round(float(row[5] or 0.0), 2),
                "frequency": str(row[6] or "Mensal"),
                "day_of_month": int(row[7] or 1),
                "payment_method": str(row[8] or "Débito Direto"),
                "last_processed_month": str(row[9] or ""),
                "active": bool(row[10]) if row[10] is not None else True,
            })
        wb.close()
        return items

    def save_recurring(self, rec_data: Dict[str, Any]) -> Dict[str, Any]:
        self.ensure_workbook()
        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb["Recorrentes"]

        rec_id = str(rec_data.get("id") or f"REC-{uuid.uuid4().hex[:6].upper()}")
        description = str(rec_data.get("description") or "Dedução Recorrente").strip()
        r_type = str(rec_data.get("type") or "Despesa")
        category = str(rec_data.get("category") or "Subscrições")
        account = str(rec_data.get("account") or "Dinheiro Geral")
        amount = round(float(rec_data.get("amount") or 0.0), 2)
        frequency = str(rec_data.get("frequency") or "Mensal")
        day_of_month = max(1, min(31, int(rec_data.get("day_of_month") or 1)))
        payment_method = str(rec_data.get("payment_method") or "Débito Direto")
        last_processed = str(rec_data.get("last_processed_month") or "")
        active = bool(rec_data.get("active", True))

        found_row = None
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(row=r, column=1).value) == rec_id:
                found_row = r
                break

        row_vals = [
            rec_id,
            description,
            r_type,
            category,
            account,
            amount,
            frequency,
            day_of_month,
            payment_method,
            last_processed,
            active,
        ]

        if found_row:
            for col_idx, val in enumerate(row_vals, start=1):
                ws.cell(row=found_row, column=col_idx, value=val)
        else:
            ws.append(row_vals)
            ws.cell(row=ws.max_row, column=6).number_format = '#,##0.00 "€"'

        wb.save(self.excel_path)
        wb.close()
        return {
            "id": rec_id,
            "description": description,
            "type": r_type,
            "category": category,
            "account": account,
            "amount": amount,
            "frequency": frequency,
            "day_of_month": day_of_month,
            "payment_method": payment_method,
            "last_processed_month": last_processed,
            "active": active,
        }

    def delete_recurring(self, rec_id: str) -> bool:
        self.ensure_workbook()
        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb["Recorrentes"]
        deleted = False
        for r in range(ws.max_row, 1, -1):
            if str(ws.cell(row=r, column=1).value) == str(rec_id):
                ws.delete_rows(r, 1)
                deleted = True
                break
        if deleted:
            wb.save(self.excel_path)
        wb.close()
        return deleted

    def process_due_recurring(self, force_all_active: bool = False) -> List[Dict[str, Any]]:
        """Verifica as deduções recorrentes ativas e lança as que já venceram neste mês."""
        now = datetime.now()
        current_month = now.strftime("%Y-%m")
        current_day = now.day

        recs = self.get_recurring()
        created_txs = []

        for rec in recs:
            if not rec.get("active", True):
                continue
            if rec.get("last_processed_month") == current_month:
                continue

            if force_all_active or current_day >= int(rec.get("day_of_month", 1)):
                day_str = f"{min(now.day, int(rec.get('day_of_month', 1))):02d}"
                tx_date = f"{current_month}-{day_str}"
                tx_id = f"TX-REC-{rec['id']}-{current_month}"

                new_tx = self.add_transaction({
                    "id": tx_id,
                    "date": tx_date,
                    "time": "09:00",
                    "type": rec["type"],
                    "account": rec["account"],
                    "to_account": "",
                    "category": rec["category"],
                    "description": f"{rec['description']} (Recorrente)",
                    "amount": rec["amount"],
                    "payment_method": rec.get("payment_method", "Débito Direto"),
                    "month_year": current_month,
                })
                created_txs.append(new_tx)
                rec["last_processed_month"] = current_month
                self.save_recurring(rec)

        return created_txs
